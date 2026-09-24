#!/usr/bin/env python3
"""화면 언어팩 ↔ 엑셀 변환 (검수용).
  python3 tools/i18n_xlsx.py export [out.xlsx]   i18n/catalog.json + locales/*.json → 검수 시트(기본 i18n/ui_review.xlsx)
  python3 tools/i18n_xlsx.py import in.xlsx      검수한 시트 → locales/*.json 갱신
시트는 언어별 열(Français·ไทย·ខ្មែរ…)이 locales/의 언어팩과 같다. 새 언어팩을 추가하면 export 때 열이 자동으로 늘어난다.
가져오기: 시트에 적힌 번역만 반영한다(빈 칸은 그 문구의 번역을 지운다). {n}·{a} 같은 자리 표시가 영어와 다르면 알려 준다.
먼저 npm run i18n 으로 카탈로그를 최신으로 만드세요."""
import json, re, sys, pathlib
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
SEP = "\u0001"
PH = re.compile(r"\{[a-z]+\}")
cat = json.load(open(ROOT / "i18n/catalog.json", encoding="utf-8"))["items"]
keyof = lambda e: e["en"] + SEP + e["ko"] if e["ambiguous"] else e["en"]
packs = []
for f in sorted((ROOT / "locales").glob("*.json")):
    d = json.load(open(f, encoding="utf-8"))
    if d.get("meta", {}).get("code"): packs.append((f, d))
packs.sort(key=lambda p: (p[1]["meta"].get("order", 99), p[1]["meta"]["code"]))
TIER = {1: "1 화면", 2: "2 분석·긴 안내"}

def export(out):
    wb = Workbook(); ws0 = wb.active; ws0.title = "사용법"
    lines = ["앱 화면 문구 검수 시트", "",
             "1. 노란 칸(언어별 열)은 제가 쓴 초안입니다. 고칠 칸만 고치세요. English와 한국어 칸은 참고용이니 고치지 마세요(열쇠로 쓰입니다).",
             "2. {n}, {a} 같은 중괄호 표시는 화면에서 숫자·이름으로 바뀌는 자리입니다. 그대로 두세요.",
             "3. 앞뒤 공백은 신경 쓰지 않아도 됩니다(원본과 같게 자동으로 맞춥니다).",
             "4. 좁은 자리(메뉴 단추, 탭)의 문구는 짧을수록 좋습니다. '넘침 주의' 열에 ▲가 있으면 원문보다 훨씬 긴 번역입니다.",
             "5. 빈 칸은 그 문구가 영어로 나온다는 뜻입니다. 2단계(분석·긴 안내) 문구는 아직 번역 전이라 비어 있습니다.",
             "6. 다 하신 뒤 이 파일을 다시 올려 주시면 언어팩에 반영합니다."]
    for r, t in enumerate(lines, 1):
        c = ws0.cell(row=r, column=1, value=t); c.font = Font(bold=(r == 1), size=13 if r == 1 else 11); c.alignment = Alignment(wrap_text=True)
    ws0.column_dimensions["A"].width = 110
    ws = wb.create_sheet("화면 문구")
    hdr = ["단계", "구획", "English", "한국어"] + [p[1]["meta"]["label"] for p in packs] + ["넘침 주의"]
    ws.append(hdr)
    head = PatternFill("solid", fgColor="D9E2F3"); yellow = PatternFill("solid", fgColor="FFF4C2"); grey = PatternFill("solid", fgColor="F2F2F2"); amber = PatternFill("solid", fgColor="FFE0B2")
    for c in ws[1]: c.font = Font(bold=True); c.fill = head; c.alignment = Alignment(wrap_text=True)
    items = sorted(cat, key=lambda e: (e["tier"],))   # 단계 순, 같은 단계는 원래 순서
    for e in items:
        vals = [p[1]["strings"].get(keyof(e), "") for p in packs]
        wide = any(v and e["words"] <= 3 and len(v.strip()) > max(14, len(e["en"].strip()) * 2.2) for v in vals)
        ws.append([TIER[e["tier"]], e["section"], e["en"], e["ko"], *vals, "▲" if wide else ""])
        r = ws.max_row
        for c in range(1, 5): ws.cell(row=r, column=c).fill = grey
        for k in range(len(packs)): ws.cell(row=r, column=5 + k).fill = yellow if vals[k] else PatternFill("solid", fgColor="FFFFFF")
        if wide: ws.cell(row=r, column=5 + len(packs)).fill = amber
    for i, w in enumerate([12, 16, 44, 36] + [40] * len(packs) + [9], 1): ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "E2"; ws.auto_filter.ref = f"A1:{get_column_letter(len(hdr))}{ws.max_row}"
    out = pathlib.Path(out or ROOT / "i18n/ui_review.xlsx"); wb.save(out)
    print(f"{out} — {len(items)}행 · 언어 {', '.join(p[1]['meta']['code'] for p in packs)}")

def do_import(path):
    ws = load_workbook(path, data_only=True)["화면 문구"]
    hdr = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(hdr)}
    by_en = {}
    for e in cat: by_en.setdefault(e["en"].strip(), []).append(e)
    labels = {p[1]["meta"]["label"]: p for p in packs}
    langcols = [(col[l], p) for l, p in labels.items() if l in col]
    changed = {p[1]["meta"]["code"]: 0 for _, p in langcols}; problems = []; unknown = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        en = (row[col["English"]] or "").strip(); ko = (row[col["한국어"]] or "").strip()
        if not en: continue
        cands = by_en.get(en, [])
        e = next((x for x in cands if x["ko"].strip() == ko), cands[0] if len(cands) == 1 else None)
        if not e: unknown += 1; continue
        key = keyof(e)
        lead = re.match(r"\s*", e["en"]).group(0); trail = re.search(r"\s*$", e["en"]).group(0)
        for ci, (f, d) in langcols:
            v = (row[ci] or "").strip() if ci < len(row) else ""
            old = d["strings"].get(key)
            if v:
                if sorted(PH.findall(e["en"])) != sorted(PH.findall(v)): problems.append(f"{d['meta']['code']} 자리 표시 불일치: {e['en'][:40]!r} → {v[:40]!r}")
                v = lead + v + trail
                if old != v: d["strings"][key] = v; changed[d["meta"]["code"]] += 1
            elif old is not None:
                del d["strings"][key]; changed[d["meta"]["code"]] += 1
    for _, (f, d) in langcols: json.dump(d, open(f, "w", encoding="utf-8"), ensure_ascii=False, indent=0)
    print("가져오기 완료 — 언어별 변경:", changed, "· 카탈로그에 없는 행", unknown)
    for p in problems[:20]: print("  !", p)

if __name__ == "__main__":
    a = sys.argv[1:]
    if not a: print(__doc__); sys.exit(1)
    if a[0] == "export": export(a[1] if len(a) > 1 else None)
    elif a[0] == "import": do_import(a[1])
    else: print(__doc__); sys.exit(1)
